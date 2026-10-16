from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime
from collections import defaultdict, Counter
import json, csv, hashlib, sys, re, os, calendar, math
from openpyxl.utils.datetime import from_excel
path,engine_path,output=sys.argv[1:4]
shared_path=sys.argv[4] if len(sys.argv)>4 else None
root=os.path.abspath(output)+os.sep
def norm(v):return str(v if v is not None else '').strip()
def amount(v):
 if isinstance(v,str):
  v=re.sub(r'\s+', '', v.replace('R$',''))
  if ',' in v:v=v.replace('.','').replace(',','.')
 if not re.fullmatch(r'-?\d+(?:\.\d+)?',str(v)):return None
 try:return int((Decimal(str(v))*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
 except:return None
def iso(v):
 if isinstance(v,(date,datetime)):return v.strftime('%Y-%m-%d')
 if isinstance(v,(int,float)) and 1<=v<2958466:
  return from_excel(int(v),w.epoch).strftime('%Y-%m-%d')
 if isinstance(v,str):
  text=v.strip();br=re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})',text);iso_match=re.fullmatch(r'(\d{4})-(\d{2})-(\d{2})(?:T.*)?',text)
  try:
   if br:return date(int(br[3]),int(br[2]),int(br[1])).isoformat()
   if iso_match:return date(*map(int,iso_match.groups())).isoformat()
  except ValueError:pass
 return None
w=load_workbook(path,data_only=True,read_only=True);rows=[];groups=defaultdict(list)
for i,raw in enumerate(w.worksheets[0].iter_rows(min_row=2,values_only=True),2):
 if not any(v is not None and v!='' for v in raw):continue
 company,appointment=norm(raw[0]).upper(),norm(raw[1]);surgery,billing=iso(raw[2]),iso(raw[11]);cents=amount(raw[13]);empty=not norm(raw[11]);errors=[]
 if len(raw)!=14:errors.append('columns')
 if not company or not appointment:errors.append('identity')
 if not surgery:errors.append('surgery_date')
 if not empty and not billing:errors.append('billing_date')
 if surgery and billing and billing<surgery:errors.append('billing_before_surgery')
 if cents is None or cents<0:errors.append('amount')
 dims={name:(norm(raw[col]).upper() if name in ['company','customerUF','hospitalUF'] else norm(raw[col])) or 'Não informado' for name,col in {'company':0,'type':3,'customer':4,'customerUF':5,'hospital':6,'hospitalUF':7,'doctor':8,'representative':10}.items()}
 row=dict(dims=dims,row=i,company=company,appointment=appointment,surgery=surgery,billing=billing,cents=cents,blankL=empty,notePresent=bool(norm(raw[12])),errors=errors)
 rows.append(row)
 if not errors:groups[(company,appointment)].append(row)
for group in groups.values():
 if len({r['surgery'] for r in group})>1:
  for r in group:r['errors'].append('conflicting_surgery_dates')
e=json.load(open(engine_path));ts={r['sourceRow']:r for r in e['records']};valid=[r for r in rows if not r['errors']];rawtotal=sum(r['cents'] or 0 for r in rows);excel=sum(r['cents'] or 0 for r in rows if r['blankL'])
assert set(ts)=={r['row'] for r in valid},'Independent valid row set differs'
fields={'company':'company','appointment':'appointment','surgeryDate':'surgery','billingDate':'billing','cents':'cents'}
assert all(all(t[k]==r[v] for k,v in fields.items()) for r in valid for t in [ts[r['row']]]),'Independent interpretation differs'
assert rawtotal==e['audit']['totalCents']
assert sum(r['cents'] for r in valid)==e['audit']['validCents']
assert sum(r['cents'] or 0 for r in rows if r['errors'])==e['audit']['quarantinedCents']
def independently_at(at,selected=None):
 gs=defaultdict(list)
 for r in valid:
  if r['surgery']<=at:gs[(r['company'],r['appointment'])].append(r)
 surgeries=[]
 for key,group in gs.items():
  pending=[r for r in group if not r['billing'] or r['billing']>at]
  billed=[r for r in group if r['billing'] and r['billing']<=at]
  status='Parcialmente faturada' if pending and billed else 'Totalmente pendente' if pending else 'Totalmente faturada'
  if selected and status!=selected:continue
  days=(date.fromisoformat(at)-date.fromisoformat(group[0]['surgery'])).days
  band=next((i for i,limit in enumerate([45,90,135,180]) if days<=limit),4) if pending else None
  surgeries.append(dict(key=key,rows=group,pending=pending,billed=billed,status=status,days=days,band=band))
 return surgeries

def summary_of(surgeries):
 pending=[s for s in surgeries if s['pending']];total=sum(r['cents'] for s in pending for r in s['pending'])
 return dict(pending=total,count=len(pending),mean=sum(s['days'] for s in pending)/len(pending) if pending else 0,partial=sum(s['status']=='Parcialmente faturada' for s in surgeries),billed=sum(s['status']=='Totalmente faturada' for s in surgeries),bands=[dict(cents=sum(r['cents'] for s in pending if s['band']==i for r in s['pending']),count=sum(s['band']==i for s in pending),mean=sum(s['days'] for s in pending if s['band']==i)/sum(s['band']==i for s in pending) if any(s['band']==i for s in pending) else 0) for i in range(5)])

def assert_summary(actual,expected):
 for field in ['pending','count','partial','billed']:assert actual[field]==expected[field],field
 assert math.isclose(actual['mean'],expected['mean'],abs_tol=1e-9)
 for a,b in zip(actual['bands'],expected['bands']):
  assert a['cents']==b['cents'] and a['count']==b['count'] and math.isclose(a['mean'],b['mean'],abs_tol=1e-9)

positions=[];trace=[]

for position in e['positions']:
 at=position['at'];gs=defaultdict(list)
 assert position['publicFinancialViewEqual'],'Public projection changed a financial view'
 for r in valid:
  if r['surgery']<=at:gs[(r['company'],r['appointment'])].append(r)
 statuses={k:('Parcialmente faturada' if any(r['billing'] and r['billing']<=at for r in group) else 'Totalmente pendente') if any(not r['billing'] or r['billing']>at for r in group) else 'Totalmente faturada' for k,group in gs.items()}
 cats=defaultdict(lambda:{'count':0,'cents':0});pending_by_status=defaultdict(int)
 for r in rows:
  included=not r['errors'] and r['surgery']<=at and (r['blankL'] or r['billing']>at)
  if r['blankL']:
   cat='excel_only_quarantine' if r['errors'] else 'excel_only_future_surgery' if r['surgery']>at else 'both_blank_pending'
  else:cat='dashboard_only_billed_after_reference' if included else 'neither'
  cats[cat]['count']+=1;cats[cat]['cents']+=r['cents'] or 0
  status=statuses.get((r['company'],r['appointment'])) if not r['errors'] and r['surgery']<=at else None
  if included:pending_by_status[status]+=r['cents']
  trace.append(dict(reference=at,sourceRow=r['row'],company=r['company'],appointment=r['appointment'],surgeryDate=r['surgery'],billingDate=r['billing'],cents=r['cents'],notePresent=r['notePresent'],excelIncluded=r['blankL'],dashboardIncluded=included,surgeryStatus=status,reason=cat,key=json.dumps([r['company'],r['appointment']],ensure_ascii=False,separators=(',',':')),engineAmountCents=ts.get(r['row'],{}).get('cents'),financialState='quarantine' if r['errors'] else 'future_surgery' if r['surgery']>at else 'pending' if included else 'billed',quarantineCodes='|'.join(r['errors'])))
 dp=sum(pending_by_status.values());bridge=excel-cats['excel_only_quarantine']['cents']-cats['excel_only_future_surgery']['cents']+cats['dashboard_only_billed_after_reference']['cents']
 assert dp==bridge==position['byStatus']['all']['summary']['pending']
 for status in ['Totalmente pendente','Parcialmente faturada','Totalmente faturada']:
  assert pending_by_status[status]==position['byStatus'][status]['summary']['pending']
 assert sum(b['cents'] for b in position['byStatus']['all']['summary']['bands'])==dp
 # Compare KPI/flow by status and independent monthly/quarterly closings.
 for label,actual in position['byStatus'].items():
  surgeries=independently_at(at,None if label=='all' else label)
  assert_summary(actual['summary'],summary_of(surgeries))
  billed_rows=[(s['key'],r) for s in surgeries for r in s['billed'] if r['billing'][:4]==at[:4]]
  assert actual['flowCents']==sum(r['cents'] for k,r in billed_rows)
  assert actual['flowCount']==len(set(k for k,r in billed_rows))
 for item in position['monthly']+position['quarterly']:assert_summary(item,summary_of(independently_at(item['at'])))
 for month in range(1,int(at[5:7])+1):
  expected=min(at,f'{at[:4]}-{month:02d}-{calendar.monthrange(int(at[:4]),month)[1]:02d}')
  assert position['monthly'][month-1]['at']==expected
 surgeries=independently_at(at)
 for dimension,ranking in position['rankings'].items():
  values=defaultdict(list)
  for s in surgeries:
   for r in s['pending']:values[r['dims'][dimension]].append((s['key'],s['days'],r['cents']))
  assert {item['name'] for item in ranking}==set(values)
  for item in ranking:
   entries=values[item['name']];cents=sum(v for k,d,v in entries);days_by_key={k:d for k,d,v in entries}
   assert item['cents']==cents and item['count']==len(days_by_key)
   assert math.isclose(item['weightedMean'],sum(d*v for k,d,v in entries)/cents if cents else 0,abs_tol=1e-8)
   assert math.isclose(item['mean'],sum(days_by_key.values())/len(days_by_key),abs_tol=1e-9)
   assert math.isclose(item['share'],cents/dp if dp else 0,abs_tol=1e-12)
 positions.append(dict(reference=at,excelBlankCents=excel,dashboardPendingCents=dp,differenceCents=dp-excel,categories=dict(cats),pendingByStatus=dict(pending_by_status),bridgeResidualCents=dp-bridge))
shared_comparison={'provided':False}
if shared_path:
 shared=json.load(open(shared_path));shared_records={r['sourceRow']:r for r in shared['records']}
 financial_fields=['company','appointment','surgeryDate','billingDate','cents']
 shared_comparison=dict(provided=True,snapshotId=shared.get('id'),snapshotAsOf=shared.get('asOf'),sameReference=shared.get('asOf')==e['asOf'],duplicateSharedRows=len(shared['records'])-len(shared_records),sameFinancialRows=len(shared['records'])==len(shared_records) and set(ts)==set(shared_records) and all(all(ts[n][f]==shared_records[n].get(f) for f in financial_fields) for n in ts),onlySourceRows=len(set(ts)-set(shared_records)),onlySharedRows=len(set(shared_records)-set(ts)),changedRows=sum(any(ts[n][f]!=shared_records[n].get(f) for f in financial_fields) for n in set(ts)&set(shared_records)))
summary=dict(sourceFile=os.path.basename(path),sourceSha256=hashlib.sha256(open(path,'rb').read()).hexdigest(),sourceRows=len(rows),validRows=len(valid),quarantineRows=len(rows)-len(valid),excelBlankCount=sum(r['blankL'] for r in rows),excelBlankCents=excel,totalCents=rawtotal,quarantineCents=e['audit']['quarantinedCents'],quarantineByCode=dict(Counter(code for r in rows for code in r['errors'])),comparedValidRows=len(ts),interpretationDifferences=0,publicFinancialViewsEqual=True,kpisBandsFlowsHistoryRankingsVerified=True,sharedComparison=shared_comparison,positions=positions)
with open(root+'row-reconciliation.csv','w') as f:
 writer=csv.DictWriter(f,fieldnames=list(trace[0]));writer.writeheader();writer.writerows(trace)
open(root+'reconciliation.json','w').write(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps(summary,ensure_ascii=False,indent=2))

for name in ['row-reconciliation.csv','reconciliation.json']:os.chmod(root+name,0o600)
